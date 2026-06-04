import json
import io
import pandas as pd

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.db.models import Avg

from .models import Station, DonneesMensuelles, BilanHydro
from .calculs import calcul_annuel_complet

MOIS_NOMS = ['Jan', 'Fév', 'Mar', 'Avr', 'Mai', 'Jun',
             'Jul', 'Aoû', 'Sep', 'Oct', 'Nov', 'Déc']


# ══════════════════════════════════════════════════════════
#  AUTHENTIFICATION
# ══════════════════════════════════════════════════════════

def login_view(request):
    if request.user.is_authenticated:
        return redirect('bilan_accueil')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            next_url = request.GET.get('next', 'accueil')
            return redirect(next_url)
        else:
            messages.error(request, "Nom d'utilisateur ou mot de passe incorrect.")

    return render(request, 'bilan_hydro/login.html')


def logout_view(request):
    logout(request)
    messages.success(request, "Vous avez été déconnecté avec succès.")
    return redirect('login')


def register_view(request):
    if request.user.is_authenticated:
        return redirect('bilan_accueil')

    if request.method == 'POST':
        username   = request.POST.get('username', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        last_name  = request.POST.get('last_name', '').strip()
        email      = request.POST.get('email', '').strip()
        password1  = request.POST.get('password1', '')
        password2  = request.POST.get('password2', '')

        # Validations
        errors = []
        if not username:
            errors.append("Le nom d'utilisateur est obligatoire.")
        elif User.objects.filter(username=username).exists():
            errors.append("Ce nom d'utilisateur est déjà pris.")
        if not password1:
            errors.append("Le mot de passe est obligatoire.")
        elif len(password1) < 6:
            errors.append("Le mot de passe doit contenir au moins 6 caractères.")
        elif password1 != password2:
            errors.append("Les deux mots de passe ne correspondent pas.")
        if email and User.objects.filter(email=email).exists():
            errors.append("Cette adresse e-mail est déjà utilisée.")

        if errors:
            for e in errors:
                messages.error(request, e)
        else:
            user = User.objects.create_user(
                username=username,
                password=password1,
                email=email,
                first_name=first_name,
                last_name=last_name,
            )
            login(request, user)
            messages.success(request,
                f"Bienvenue {first_name or username} ! Votre compte a été créé avec succès.")
            return redirect('bilan_accueil')

    return render(request, 'bilan_hydro/register.html')


# ══════════════════════════════════════════════════════════
#  VUES PRINCIPALES (protégées)
# ══════════════════════════════════════════════════════════

@login_required
def accueil(request):  # bilan_accueil
    stations = Station.objects.all()
    return render(request, 'bilan_hydro/accueil.html', {'stations': stations, 'module': 'bilan'})


@login_required
def import_donnees(request):
    """Page d'import de données Excel via interface web."""
    stations = Station.objects.all()

    if request.method == 'POST':
        fichier = request.FILES.get('fichier')
        nom_station = request.POST.get('nom_station', '').strip()
        latitude = request.POST.get('latitude', '').strip()
        remplacement = request.POST.get('remplacement') == 'on'

        # Validations
        if not fichier:
            messages.error(request, "Veuillez sélectionner un fichier Excel.")
            return render(request, 'bilan_hydro/import.html', {'stations': stations})

        if not fichier.name.endswith(('.xlsx', '.xls')):
            messages.error(request, "Format non supporté. Utilisez un fichier .xlsx ou .xls")
            return render(request, 'bilan_hydro/import.html', {'stations': stations})

        if not nom_station:
            messages.error(request, "Veuillez saisir le nom de la station.")
            return render(request, 'bilan_hydro/import.html', {'stations': stations})

        try:
            latitude = float(latitude)
        except (ValueError, TypeError):
            messages.error(request, "Latitude invalide. Entrez un nombre (ex: -18.9)")
            return render(request, 'bilan_hydro/import.html', {'stations': stations})

        # Lecture du fichier
        try:
            contenu = fichier.read()
            df = pd.read_excel(io.BytesIO(contenu), skiprows=2, header=0)
            df.columns = ['Date', 'Temperature', 'Precipitation', 'Humidite', 'Insolation']
            df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
            df = df.dropna(subset=['Date'])
            for col in ['Temperature', 'Precipitation', 'Humidite', 'Insolation']:
                df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)
        except Exception as e:
            messages.error(request, f"Erreur de lecture du fichier : {e}")
            return render(request, 'bilan_hydro/import.html', {'stations': stations})

        if df.empty:
            messages.error(request, "Le fichier ne contient aucune donnée valide.")
            return render(request, 'bilan_hydro/import.html', {'stations': stations})

        # Agrégation mensuelle
        df['annee'] = df['Date'].dt.year
        df['mois'] = df['Date'].dt.month
        monthly = df.groupby(['annee', 'mois']).agg(
            temperature_moy=('Temperature', 'mean'),
            precipitation=('Precipitation', 'sum'),
            humidite_moy=('Humidite', 'mean'),
            insolation=('Insolation', 'sum')
        ).reset_index()

        # Créer ou mettre à jour la station
        station, created = Station.objects.get_or_create(
            nom=nom_station,
            defaults={'latitude': latitude}
        )
        if not created:
            station.latitude = latitude
            station.save()

        # Supprimer les anciennes données si demandé
        if remplacement:
            DonneesMensuelles.objects.filter(station=station).delete()
            BilanHydro.objects.filter(station=station).delete()

        # Insertion des données mensuelles
        nb_crees = 0
        nb_maj = 0
        for _, row in monthly.iterrows():
            obj, was_created = DonneesMensuelles.objects.update_or_create(
                station=station,
                annee=int(row['annee']),
                mois=int(row['mois']),
                defaults={
                    'temperature_moy': round(float(row['temperature_moy']), 4),
                    'precipitation': round(float(row['precipitation']), 4),
                    'humidite_moy': round(float(row['humidite_moy']), 4),
                    'insolation': round(float(row['insolation']), 4),
                }
            )
            if was_created:
                nb_crees += 1
            else:
                nb_maj += 1

        annee_min = int(monthly['annee'].min())
        annee_max = int(monthly['annee'].max())
        nb_lignes = len(df)

        messages.success(request,
            f"✅ Import réussi ! Station «{nom_station}» — "
            f"{nb_crees} mois créés, {nb_maj} mis à jour — "
            f"Période : {annee_min}–{annee_max} ({nb_lignes} lignes brutes traitées)."
        )
        return redirect('calcul_bilan', station_id=station.id)

    return render(request, 'bilan_hydro/import.html', {'stations': stations})


@login_required
def supprimer_station(request, station_id):
    """Supprime une station et toutes ses données."""
    station = get_object_or_404(Station, pk=station_id)
    if request.method == 'POST':
        nom = station.nom
        station.delete()
        messages.success(request, f"Station «{nom}» supprimée avec succès.")
        return redirect('bilan_accueil')
    return render(request, 'bilan_hydro/confirmer_suppression.html', {'station': station})


@login_required
def station_detail(request, station_id):
    station = get_object_or_404(Station, pk=station_id)
    annees = DonneesMensuelles.objects.filter(station=station).values_list(
        'annee', flat=True).distinct().order_by('annee')
    return render(request, 'bilan_hydro/station_detail.html', {
        'station': station,
        'annees': list(annees),
    })


@login_required
def calcul_bilan(request, station_id):
    station = get_object_or_404(Station, pk=station_id)
    annees_dispo = list(
        DonneesMensuelles.objects.filter(station=station)
        .values_list('annee', flat=True).distinct().order_by('annee')
    )

    if request.method == 'POST':
        methode = request.POST.get('methode', 'thornthwaite')
        mode = request.POST.get('mode', 'annee')  # 'annee' ou 'periode'
        rfu_max = float(request.POST.get('rfu_max', 100))

        if mode == 'annee':
            annee = int(request.POST.get('annee', annees_dispo[-1]))
            annees_calc = [annee]
        else:
            an_debut = int(request.POST.get('annee_debut', annees_dispo[0]))
            an_fin = int(request.POST.get('annee_fin', annees_dispo[-1]))
            annees_calc = list(range(an_debut, an_fin + 1))

        resultats_annuels = []
        for annee in annees_calc:
            donnees = DonneesMensuelles.objects.filter(
                station=station, annee=annee
            ).order_by('mois')

            if donnees.count() < 12:
                continue

            d = {
                'annee': annee,
                'temperatures': [r.temperature_moy for r in donnees],
                'precipitations': [r.precipitation for r in donnees],
                'insolations': [r.insolation for r in donnees],
                'humidites': [r.humidite_moy for r in donnees],
            }

            res = calcul_annuel_complet(d, station.latitude, methode, rfu_max)

            # Sauvegarder en BD
            for i in range(12):
                BilanHydro.objects.update_or_create(
                    station=station, methode=methode, annee=annee, mois=i + 1,
                    defaults={
                        'etp': res['etp'][i],
                        'etr': res['etr'][i],
                        'rfu': res['rfu'][i],
                        'da': res['da'][i],
                        'variation_stock': res['variation_stock'][i],
                        'excedent': res['excedent'][i],
                    }
                )

            resultats_annuels.append({
                'annee': annee,
                'mois': MOIS_NOMS,
                'precipitation': d['precipitations'],
                'temperature': [round(t, 2) for t in d['temperatures']],
                'etp': res['etp'],
                'etr': res['etr'],
                'rfu': res['rfu'],
                'da': res['da'],
                'variation_stock': res['variation_stock'],
                'excedent': res['excedent'],
            })

        return render(request, 'bilan_hydro/resultats.html', {
            'station': station,
            'methode': methode,
            'methode_label': 'Thornthwaite' if methode == 'thornthwaite' else 'Turc',
            'resultats': resultats_annuels,
            'resultats_json': json.dumps(resultats_annuels),
            'rfu_max': rfu_max,
            'mode': mode,
        })

    return render(request, 'bilan_hydro/calcul_form.html', {
        'station': station,
        'annees': annees_dispo,
    })


@login_required
def historique(request, station_id):
    station = get_object_or_404(Station, pk=station_id)
    methode = request.GET.get('methode', 'thornthwaite')
    bilans = BilanHydro.objects.filter(station=station, methode=methode).order_by('annee', 'mois')

    # Moyennes mensuelles interannuelles
    moyennes = []
    for m in range(1, 13):
        b = bilans.filter(mois=m).aggregate(
            etp=Avg('etp'), etr=Avg('etr'), da=Avg('da'),
            rfu=Avg('rfu'), excedent=Avg('excedent')
        )
        moyennes.append({
            'mois': MOIS_NOMS[m - 1],
            'etp': round(b['etp'] or 0, 2),
            'etr': round(b['etr'] or 0, 2),
            'da': round(b['da'] or 0, 2),
            'rfu': round(b['rfu'] or 0, 2),
            'excedent': round(b['excedent'] or 0, 2),
        })

    return render(request, 'bilan_hydro/historique.html', {
        'station': station,
        'methode': methode,
        'methode_label': 'Thornthwaite' if methode == 'thornthwaite' else 'Turc',
        'moyennes': moyennes,
        'moyennes_json': json.dumps(moyennes),
        'bilans': bilans,
    })


@login_required
def api_donnees(request, station_id):
    """API JSON pour les graphiques dynamiques."""
    station = get_object_or_404(Station, pk=station_id)
    annee = int(request.GET.get('annee', 0))
    methode = request.GET.get('methode', 'thornthwaite')

    bilans = BilanHydro.objects.filter(
        station=station, methode=methode, annee=annee
    ).order_by('mois')

    data = {
        'mois': MOIS_NOMS,
        'etp': [b.etp for b in bilans],
        'etr': [b.etr for b in bilans],
        'rfu': [b.rfu for b in bilans],
        'da': [b.da for b in bilans],
        'variation_stock': [b.variation_stock for b in bilans],
        'excedent': [b.excedent for b in bilans],
    }
    return JsonResponse(data)


@login_required
def export_excel(request, station_id):
    """Exporte les résultats du bilan en fichier Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import (PatternFill, Font, Alignment,
                                  Border, Side, GradientFill)
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference, Series
    from openpyxl.chart.series import SeriesLabel
    from django.http import HttpResponse

    station = get_object_or_404(Station, pk=station_id)
    methode = request.GET.get('methode', 'thornthwaite')
    annees_param = request.GET.get('annees', '')

    # Années à exporter
    if annees_param:
        annees = [int(a) for a in annees_param.split(',') if a.strip().isdigit()]
    else:
        annees = list(
            BilanHydro.objects.filter(station=station, methode=methode)
            .values_list('annee', flat=True).distinct().order_by('annee')
        )

    if not annees:
        from django.contrib import messages
        messages.warning(request, "Aucun bilan calculé à exporter pour cette méthode.")
        return redirect('calcul_bilan', station_id=station_id)

    wb = openpyxl.Workbook()
    methode_label = 'Thornthwaite' if methode == 'thornthwaite' else 'Turc'

    # Styles
    HDR = PatternFill("solid", fgColor="1A5276")
    HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
    SUB_FILL = PatternFill("solid", fgColor="D6EAF8")
    SUB_FONT = Font(bold=True, size=10)
    TOT_FILL = PatternFill("solid", fgColor="F0F3F4")
    TOT_FONT = Font(bold=True, italic=True)
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='AAAAAA')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    COLOR_MAP = {
        'P':   'D6EAF8', 'T': 'FDFEFE',
        'ETP': 'E8DAEF', 'ETR': 'D5F5E3',
        'RFU': 'D6EAF8', 'DA':  'FADBD8',
        'DS':  'FEF9E7', 'EXC': 'D1F2EB',
    }

    # ── Feuille récapitulative multi-années ──────────────────────────────
    ws_recap = wb.active
    ws_recap.title = f"Récapitulatif {methode_label}"

    HEADERS = ['Année', 'Mois', 'P (mm)', 'T (°C)',
               'ETP (mm)', 'ETR (mm)', 'RFU (mm)',
               'DA (mm)', 'ΔS (mm)', 'EXC (mm)']
    COL_KEYS = ['annee', 'mois', 'precipitation', 'temperature',
                'etp', 'etr', 'rfu', 'da', 'variation_stock', 'excedent']
    COL_COLORS = [None, None,
                  COLOR_MAP['P'], COLOR_MAP['T'],
                  COLOR_MAP['ETP'], COLOR_MAP['ETR'], COLOR_MAP['RFU'],
                  COLOR_MAP['DA'], COLOR_MAP['DS'], COLOR_MAP['EXC']]

    # En-tête station
    ws_recap.merge_cells('A1:J1')
    c = ws_recap['A1']
    c.value = f"Bilan Hydrique — Station {station.nom} — Méthode {methode_label} — Lat. {station.latitude}°"
    c.font = Font(bold=True, size=13, color="1A5276")
    c.alignment = center
    ws_recap.row_dimensions[1].height = 22

    # En-têtes colonnes
    for col, (h, clr) in enumerate(zip(HEADERS, COL_COLORS), 1):
        cell = ws_recap.cell(row=2, column=col, value=h)
        cell.fill = HDR
        cell.font = HDR_FONT
        cell.alignment = center
        cell.border = border

    row_idx = 3
    totaux_rows = []  # pour les lignes TOTAL

    for annee in annees:
        donnees_m = DonneesMensuelles.objects.filter(
            station=station, annee=annee).order_by('mois')
        bilans_m = BilanHydro.objects.filter(
            station=station, methode=methode, annee=annee).order_by('mois')

        dm = {d.mois: d for d in donnees_m}
        bm = {b.mois: b for b in bilans_m}

        sums = {k: 0.0 for k in ['precipitation', 'etp', 'etr', 'da', 'excedent']}

        for mois_n in range(1, 13):
            d = dm.get(mois_n)
            b = bm.get(mois_n)
            row_data = {
                'annee': annee,
                'mois': MOIS_NOMS[mois_n - 1],
                'precipitation': round(d.precipitation, 1) if d else '—',
                'temperature': round(d.temperature_moy, 1) if d else '—',
                'etp': round(b.etp, 1) if b else '—',
                'etr': round(b.etr, 1) if b else '—',
                'rfu': round(b.rfu, 1) if b else '—',
                'da': round(b.da, 1) if b else '—',
                'variation_stock': round(b.variation_stock, 1) if b else '—',
                'excedent': round(b.excedent, 1) if b else '—',
            }
            for k in sums:
                v = row_data.get(k)
                if isinstance(v, (int, float)):
                    sums[k] += v

            for col, key in enumerate(COL_KEYS, 1):
                cell = ws_recap.cell(row=row_idx, column=col, value=row_data[key])
                cell.alignment = center
                cell.border = border
                clr = COL_COLORS[col - 1]
                if clr:
                    cell.fill = PatternFill("solid", fgColor=clr)
            row_idx += 1

        # Ligne TOTAL
        tot_row = row_idx
        totaux_rows.append(tot_row)
        tot_data = ['TOTAL', '—',
                    round(sums['precipitation'], 1), '—',
                    round(sums['etp'], 1), round(sums['etr'], 1), '—',
                    round(sums['da'], 1), '—', round(sums['excedent'], 1)]
        for col, val in enumerate(tot_data, 1):
            cell = ws_recap.cell(row=tot_row, column=col, value=val)
            cell.fill = TOT_FILL
            cell.font = TOT_FONT
            cell.alignment = center
            cell.border = border
        row_idx += 1

    # Largeurs colonnes recap
    widths = [8, 6, 9, 8, 9, 9, 9, 9, 9, 9]
    for i, w in enumerate(widths, 1):
        ws_recap.column_dimensions[get_column_letter(i)].width = w

    # ── Feuilles individuelles par année ─────────────────────────────────
    for annee in annees:
        ws = wb.create_sheet(title=str(annee))

        donnees_m = DonneesMensuelles.objects.filter(
            station=station, annee=annee).order_by('mois')
        bilans_m = BilanHydro.objects.filter(
            station=station, methode=methode, annee=annee).order_by('mois')
        dm = {d.mois: d for d in donnees_m}
        bm = {b.mois: b for b in bilans_m}

        # Titre
        ws.merge_cells('A1:J1')
        c = ws['A1']
        c.value = f"{station.nom} — {annee} — {methode_label}"
        c.font = Font(bold=True, size=12, color="1A5276")
        c.alignment = center
        ws.row_dimensions[1].height = 20

        for col, (h, clr) in enumerate(zip(HEADERS, COL_COLORS), 1):
            cell = ws.cell(row=2, column=col, value=h)
            cell.fill = HDR
            cell.font = HDR_FONT
            cell.alignment = center
            cell.border = border

        etp_vals, etr_vals, p_vals, da_vals, rfu_vals, exc_vals = [], [], [], [], [], []
        sums = {k: 0.0 for k in ['precipitation', 'etp', 'etr', 'da', 'excedent']}

        for mois_n in range(1, 13):
            d = dm.get(mois_n)
            b = bm.get(mois_n)
            row_data = {
                'annee': annee,
                'mois': MOIS_NOMS[mois_n - 1],
                'precipitation': round(d.precipitation, 1) if d else 0,
                'temperature': round(d.temperature_moy, 1) if d else 0,
                'etp': round(b.etp, 1) if b else 0,
                'etr': round(b.etr, 1) if b else 0,
                'rfu': round(b.rfu, 1) if b else 0,
                'da': round(b.da, 1) if b else 0,
                'variation_stock': round(b.variation_stock, 1) if b else 0,
                'excedent': round(b.excedent, 1) if b else 0,
            }
            p_vals.append(row_data['precipitation'])
            etp_vals.append(row_data['etp'])
            etr_vals.append(row_data['etr'])
            da_vals.append(row_data['da'])
            rfu_vals.append(row_data['rfu'])
            exc_vals.append(row_data['excedent'])
            for k in sums:
                sums[k] += row_data.get(k, 0)

            r = mois_n + 2
            for col, key in enumerate(COL_KEYS, 1):
                cell = ws.cell(row=r, column=col, value=row_data[key])
                cell.alignment = center
                cell.border = border
                clr = COL_COLORS[col - 1]
                if clr:
                    cell.fill = PatternFill("solid", fgColor=clr)

        # Ligne TOTAL
        tr = 15
        tot_data = ['TOTAL', '—',
                    round(sums['precipitation'], 1), '—',
                    round(sums['etp'], 1), round(sums['etr'], 1), '—',
                    round(sums['da'], 1), '—', round(sums['excedent'], 1)]
        for col, val in enumerate(tot_data, 1):
            cell = ws.cell(row=tr, column=col, value=val)
            cell.fill = TOT_FILL
            cell.font = TOT_FONT
            cell.alignment = center
            cell.border = border

        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # ── Graphique Bilan (barres P + lignes ETP/ETR/DA) ──────────────
        chart1 = BarChart()
        chart1.type = "col"
        chart1.title = f"Bilan Hydrique {annee}"
        chart1.y_axis.title = "mm"
        chart1.x_axis.title = "Mois"
        chart1.style = 10
        chart1.width = 18
        chart1.height = 12

        cats = Reference(ws, min_col=2, min_row=3, max_row=14)

        data_p = Reference(ws, min_col=3, min_row=2, max_row=14)
        s_p = BarChart()
        s_p.type = "col"
        bar_p = Reference(ws, min_col=3, min_row=2, max_row=14)
        chart1.add_data(bar_p, titles_from_data=True)
        chart1.set_categories(cats)

        # Lignes ETP, ETR, DA
        for col_n, title, color in [(5, 'ETP', '8E44AD'), (6, 'ETR', '27AE60'), (8, 'DA', 'C0392B')]:
            line_chart = LineChart()
            line_data = Reference(ws, min_col=col_n, min_row=2, max_row=14)
            line_chart.add_data(line_data, titles_from_data=True)
            line_chart.set_categories(cats)
            chart1 += line_chart

        ws.add_chart(chart1, "A17")

        # ── Graphique RFU / Excédent ─────────────────────────────────────
        chart2 = BarChart()
        chart2.type = "col"
        chart2.title = f"RFU & Excédent {annee}"
        chart2.y_axis.title = "mm"
        chart2.x_axis.title = "Mois"
        chart2.style = 10
        chart2.width = 18
        chart2.height = 12

        rfu_ref = Reference(ws, min_col=7, min_row=2, max_row=14)
        chart2.add_data(rfu_ref, titles_from_data=True)
        chart2.set_categories(cats)

        exc_line = LineChart()
        exc_ref = Reference(ws, min_col=10, min_row=2, max_row=14)
        exc_line.add_data(exc_ref, titles_from_data=True)
        exc_line.set_categories(cats)
        chart2 += exc_line

        ws.add_chart(chart2, "K17")

    # ── Feuille comparaison ETP (si les 2 méthodes sont dispo) ──────────
    methode_autre = 'turc' if methode == 'thornthwaite' else 'thornthwaite'
    autre_label = 'Turc' if methode == 'thornthwaite' else 'Thornthwaite'
    bilans_autre = BilanHydro.objects.filter(
        station=station, methode=methode_autre,
        annee__in=annees
    )
    if bilans_autre.exists():
        ws_cmp = wb.create_sheet(title="Comparaison ETP")

        ws_cmp.merge_cells('A1:N1')
        c = ws_cmp['A1']
        c.value = (f"Comparaison ETP — {station.nom} — "
                   f"{methode_label} vs {autre_label}")
        c.font = Font(bold=True, size=13, color="1A5276")
        c.alignment = center

        # En-têtes : Année | Jan…Déc | Total
        ws_cmp.cell(row=2, column=1, value='Année/Méthode').fill = HDR
        ws_cmp.cell(row=2, column=1).font = HDR_FONT
        for i, m in enumerate(MOIS_NOMS, 2):
            ws_cmp.cell(row=2, column=i, value=m).fill = HDR
            ws_cmp.cell(row=2, column=i).font = HDR_FONT
            ws_cmp.cell(row=2, column=i).alignment = center
        ws_cmp.cell(row=2, column=14, value='Total').fill = HDR
        ws_cmp.cell(row=2, column=14).font = HDR_FONT

        row_c = 3
        for annee in annees:
            b1 = list(BilanHydro.objects.filter(
                station=station, methode=methode, annee=annee
            ).order_by('mois').values_list('etp', flat=True))
            b2 = list(BilanHydro.objects.filter(
                station=station, methode=methode_autre, annee=annee
            ).order_by('mois').values_list('etp', flat=True))

            if not b1 or not b2:
                continue

            # Ligne méthode principale
            ws_cmp.cell(row=row_c, column=1,
                        value=f"{annee} — {methode_label}").fill = PatternFill("solid", fgColor="E8DAEF")
            ws_cmp.cell(row=row_c, column=1).font = Font(bold=True)
            for i, v in enumerate(b1, 2):
                ws_cmp.cell(row=row_c, column=i, value=round(v, 1)).alignment = center
            ws_cmp.cell(row=row_c, column=14, value=round(sum(b1), 1)).font = Font(bold=True)

            # Ligne méthode autre
            ws_cmp.cell(row=row_c + 1, column=1,
                        value=f"{annee} — {autre_label}").fill = PatternFill("solid", fgColor="D5F5E3")
            ws_cmp.cell(row=row_c + 1, column=1).font = Font(bold=True)
            for i, v in enumerate(b2, 2):
                ws_cmp.cell(row=row_c + 1, column=i, value=round(v, 1)).alignment = center
            ws_cmp.cell(row=row_c + 1, column=14, value=round(sum(b2), 1)).font = Font(bold=True)

            # Ligne écart
            ws_cmp.cell(row=row_c + 2, column=1, value=f"Écart").fill = PatternFill("solid", fgColor="FDEBD0")
            for i, (v1, v2) in enumerate(zip(b1, b2), 2):
                diff = round(v1 - v2, 1)
                c = ws_cmp.cell(row=row_c + 2, column=i, value=diff)
                c.alignment = center
                c.font = Font(color='C0392B' if diff < 0 else '27AE60')
            ws_cmp.cell(row=row_c + 2, column=14, value=round(sum(b1) - sum(b2), 1))

            row_c += 3

        ws_cmp.column_dimensions['A'].width = 28
        for i in range(2, 15):
            ws_cmp.column_dimensions[get_column_letter(i)].width = 8

    # ── Réponse HTTP ─────────────────────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = (f"bilan_{station.nom}_{methode_label}_"
                f"{annees[0]}-{annees[-1]}.xlsx")
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def api_comparaison_etp(request, station_id):
    """Retourne les ETP des 2 méthodes pour une même année (comparaison graphique)."""
    station = get_object_or_404(Station, pk=station_id)
    annee = int(request.GET.get('annee', 0))

    def get_etp(methode):
        bilans = BilanHydro.objects.filter(
            station=station, methode=methode, annee=annee
        ).order_by('mois')
        vals = [b.etp for b in bilans]
        return vals if len(vals) == 12 else None

    etp_thorn = get_etp('thornthwaite')
    etp_turc = get_etp('turc')

    return JsonResponse({
        'mois': MOIS_NOMS,
        'thornthwaite': etp_thorn,
        'turc': etp_turc,
        'has_thorn': etp_thorn is not None,
        'has_turc': etp_turc is not None,
        'comparaison_dispo': etp_thorn is not None and etp_turc is not None,
    })
